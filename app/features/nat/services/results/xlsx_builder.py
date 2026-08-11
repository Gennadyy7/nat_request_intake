import asyncio
from collections.abc import Sequence
from pathlib import Path
import tempfile

from openpyxl import Workbook

from app.features.nat.services.results.tabular_loader import sanitize_sheet_title


async def build_results_xlsx_file(
    *,
    sheets: Sequence[tuple[str, list[list[str]]]],
) -> Path:
    with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as handle:
        destination = Path(handle.name)

    def _write() -> None:
        workbook = Workbook()
        default_sheet = workbook.active
        if default_sheet is not None:
            workbook.remove(default_sheet)
        used_titles: set[str] = set()
        for title, rows in sheets:
            sheet_title = sanitize_sheet_title(title, used_titles=used_titles)
            worksheet = workbook.create_sheet(title=sheet_title)
            for row_index, row in enumerate(rows, start=1):
                for column_index, value in enumerate(row, start=1):
                    cell = worksheet.cell(
                        row=row_index,
                        column=column_index,
                        value=value,
                    )
                    cell.number_format = '@'
        workbook.save(destination)

    try:
        await asyncio.to_thread(_write)
    except Exception:
        destination.unlink(missing_ok=True)
        raise
    return destination
