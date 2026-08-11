import csv
import io
from pathlib import Path

from openpyxl import load_workbook

from app.features.nat.constants import AllowedFileExtension
from app.features.nat.services.parsing.file_parser import resolve_extension

CSV_DELIMITER = ';'
CSV_ENCODING = 'utf-8-sig'
EXCEL_MAX_SHEET_NAME_LENGTH = 31


def load_csv_rows(path: Path) -> list[list[str]]:
    text = path.read_text(encoding=CSV_ENCODING)
    reader = csv.reader(io.StringIO(text), delimiter=CSV_DELIMITER)
    return [[str(cell) for cell in row] for row in reader]


def load_application_rows(path: Path, *, original_filename: str) -> list[list[str]]:
    extension = resolve_extension(original_filename)
    if extension == AllowedFileExtension.XLSX:
        return _load_xlsx_rows(path)
    return load_csv_rows(path)


def merge_csv_parts(paths: list[Path]) -> list[list[str]]:
    if not paths:
        return []
    merged: list[list[str]] = []
    header_taken = False
    for path in paths:
        rows = load_csv_rows(path)
        if not rows:
            continue
        if not header_taken:
            merged.extend(rows)
            header_taken = True
            continue
        merged.extend(rows[1:])
    return merged


def sanitize_sheet_title(title: str, *, used_titles: set[str]) -> str:
    normalized = title[:EXCEL_MAX_SHEET_NAME_LENGTH]
    if normalized not in used_titles:
        used_titles.add(normalized)
        return normalized

    suffix_index = 2
    while True:
        suffix = f'_{suffix_index}'
        base_length = EXCEL_MAX_SHEET_NAME_LENGTH - len(suffix)
        candidate = f'{title[:base_length]}{suffix}'
        if candidate not in used_titles:
            used_titles.add(candidate)
            return candidate
        suffix_index += 1


def _load_xlsx_rows(path: Path) -> list[list[str]]:
    workbook = load_workbook(filename=path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        if sheet is None:
            return []
        rows: list[list[str]] = []
        for data_row in sheet.iter_rows(values_only=True):
            if data_row is None:
                continue
            rows.append([_cell_to_str(value) for value in data_row])
        return rows
    finally:
        workbook.close()


def _cell_to_str(value: object) -> str:
    if value is None:
        return ''
    return str(value).strip()
