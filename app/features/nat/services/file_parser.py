import csv
import io
from pathlib import PurePath

from openpyxl import load_workbook

from app.core.config import Settings
from app.features.nat.constants import AllowedFileExtension
from app.features.nat.services.parsed_file import ParsedFile, ParsedRow


def parse_file_content(
    content: bytes,
    extension: AllowedFileExtension,
    settings: Settings,
) -> ParsedFile:
    if extension == AllowedFileExtension.XLSX:
        return _parse_xlsx(content)
    return _parse_delimited(content, settings.NAT_OUTPUT_FIELD_SEPARATOR)


def _build_parsed_row(headers: list[str], row_values: list[str]) -> ParsedRow:
    return ParsedRow(
        values={
            headers[index]: row_values[index].strip() if index < len(row_values) else ''
            for index in range(len(headers))
        },
        raw_field_count=len(row_values),
    )


def _parse_delimited(content: bytes, delimiter: str) -> ParsedFile:
    text = content.decode('utf-8-sig')
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    table = [row for row in reader if any(cell.strip() for cell in row)]
    if not table:
        return ParsedFile(headers=[], rows=[])

    headers = [cell.strip() for cell in table[0]]
    rows = [_build_parsed_row(headers, row_values) for row_values in table[1:]]
    return ParsedFile(headers=headers, rows=rows)


def _parse_xlsx(content: bytes) -> ParsedFile:
    workbook = load_workbook(
        filename=io.BytesIO(content), read_only=True, data_only=True
    )
    try:
        sheet = workbook.active
        if sheet is None:
            return ParsedFile(headers=[], rows=[])
        row_iter = sheet.iter_rows(values_only=True)
        header_row = next(row_iter, None)
        if header_row is None:
            return ParsedFile(headers=[], rows=[])

        headers = [_cell_to_str(value) for value in header_row]
        rows: list[ParsedRow] = []
        for data_row in row_iter:
            if data_row is None or not any(
                value is not None and str(value).strip() for value in data_row
            ):
                continue
            row_values = [_cell_to_str(value) for value in data_row]
            rows.append(_build_parsed_row(headers, row_values))
        return ParsedFile(headers=headers, rows=rows)
    finally:
        workbook.close()


def _cell_to_str(value: object) -> str:
    if value is None:
        return ''
    return str(value).strip()


def resolve_extension(filename: str) -> AllowedFileExtension | None:
    suffix = PurePath(filename).suffix.lower()
    try:
        return AllowedFileExtension(suffix)
    except ValueError:
        return None
